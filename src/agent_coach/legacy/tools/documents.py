"""文档读取：简历 / 工作日志 / 自填文本 / 各种文件 → 纯文本。

支持类型（按扩展名）：
  .txt .md .markdown .log .json .csv   —— 纯文本直读
  .pdf                                —— pypdf
  .docx .docm                         —— 标准库 zipfile+xml（无第三方依赖）
  .pptx .pptm                         —— 标准库 zipfile+xml（无第三方依赖）
  .xlsx .xlsm                         —— openpyxl（纯 Python）
其余类型抛 UnsupportedDocumentError（调用方跳过并提示，不中断整批）。
"""

from __future__ import annotations

import json as json_mod
import re
import zipfile
import posixpath
from pathlib import Path
from xml.etree import ElementTree as ET

from pypdf import PdfReader

from .text_codec import decode_text_bytes

SUPPORTED_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".log", ".json", ".csv", ".pdf", ".docx", ".docm", ".pptx", ".pptm", ".xlsx", ".xlsm"})


class UnsupportedDocumentError(ValueError):
    pass


def _local(tag: str) -> str:
    """去掉 XML 命名空间前缀，取本地名。"""
    return tag.rsplit("}", 1)[-1]


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(str(path)) as z:
        if "word/document.xml" not in z.namelist():
            raise UnsupportedDocumentError(f"不是有效的 Word 文档：{path.name}")
        root = ET.fromstring(z.read("word/document.xml"))
    out: list[str] = []
    for node in root.iter():
        name = _local(node.tag)
        if name == "p":
            out.append("\n")
        elif name == "t" and node.text and node.text.strip():
            out.append(node.text)
    return "".join(out).strip()


def _pptx_text(path: Path) -> str:
    with zipfile.ZipFile(str(path)) as z:
        try:
            presentation = ET.fromstring(z.read("ppt/presentation.xml"))
            relationships = ET.fromstring(z.read("ppt/_rels/presentation.xml.rels"))
            targets = {r.attrib["Id"]: r.attrib["Target"] for r in relationships
                       if r.attrib.get("TargetMode") != "External" and r.attrib.get("Type", "").endswith("/slide")}
            names = []
            for node in presentation.iter():
                if _local(node.tag) != "sldId":
                    continue
                relation = next((v for k, v in node.attrib.items() if k.endswith("}id")), None)
                target = targets.get(relation)
                if target is None:
                    raise ValueError("Missing slide relationship")
                name = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("ppt", target))
                if not name.startswith("ppt/slides/"):
                    raise ValueError("Invalid slide target")
                names.append(name)
        except (KeyError, ValueError, ET.ParseError) as exc:
            raise UnsupportedDocumentError(f"PPT 页面关系无效：{path.name}") from exc
        if not names:
            raise UnsupportedDocumentError(f"不是有效的 PPT 文档：{path.name}")
        slides: list[str] = []
        for n in names:
            root = ET.fromstring(z.read(n))
            texts = [
                node.text.strip()
                for node in root.iter()
                if _local(node.tag) == "t" and node.text and node.text.strip()
            ]
            if texts:
                slides.append("[Slide] " + " / ".join(texts))
    return "\n".join(slides).strip()


def read_document(path: str | Path) -> str:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"文档不存在：{path}")

    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown", ".log"}:
        return decode_text_bytes(path.read_bytes()).strip()
    if suffix in {".json"}:
        text = decode_text_bytes(path.read_bytes()).strip()
        try:
            return json_mod.dumps(json_mod.loads(text), ensure_ascii=False, indent=2)
        except Exception:  # noqa: BLE001 - 不是合法 JSON 就当普通文本
            return text
    if suffix in {".csv"}:
        return decode_text_bytes(path.read_bytes()).strip()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()
    if suffix in {".docx", ".docm"}:
        return _docx_text(path)
    if suffix in {".pptx", ".pptm"}:
        return _pptx_text(path)
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        rows: list[str] = []
        for ws in wb.worksheets:
            if rows:
                rows.append(f"--- Sheet: {ws.title} ---")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    rows.append(" | ".join(vals))
        wb.close()
        return "\n".join(rows).strip()
    raise UnsupportedDocumentError(
        f"暂不支持 {suffix} 文件（{path.name}），请转为 PDF / Word / Excel / PPT / 纯文本后重试"
    )
